"""Shared helpers for the two kw.beijing.gov.cn award ingests (oxjob #1451):

    beijing_nsf_to_s3.py    -> Beijing Municipal Natural Science Foundation (F4320334977)
    beijing_bmstc_to_s3.py  -> Beijing Municipal Science and Technology Commission (F4320325902)

kw.beijing.gov.cn is ONE portal publishing for TWO funders (how-to §2.3.2): the
北京市自然科学基金委员会办公室 (基金办) rosters go to the NSF ingest; the 市科委
(BMSTC) plan-project 立项公开清单 go to the BMSTC ingest. Each runner lists its own
source documents explicitly, so nothing is blanket-assigned.

Fetching: live kw.beijing.gov.cn first; documents that only survive on the
Internet Archive are fetched as raw `id_` captures (NSSFC 475 precedent).
Parsing: .docx (python-docx), legacy .doc (Word COM via cn_provincial.extras),
.xlsx (pandas), inline HTML tables (pandas.read_html).
"""

from __future__ import annotations

import calendar
import hashlib
import io
import re
import sys
import time
from pathlib import Path
from typing import Any

import requests

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                         "(KHTML, like Gecko) Chrome/120 Safari/537.36 openalex-walden"}
RETRIES = 4


def log(msg: str) -> None:
    print(msg, flush=True)


CACHE_DIR: Path | None = None  # set by the runner; checkpoint so re-runs skip fetched docs


def fetch(url: str, wayback_ts: str | None = None, pause: float = 1.0, cache: bool = True) -> bytes:
    """GET a document. With wayback_ts, fetch the raw Internet Archive capture.

    Documents are checkpointed in CACHE_DIR (how-to §1.2 item 3), so a re-run after a
    Wayback throttle resumes instead of refetching. Raises on persistent failure -- a
    missing source document must fail the run loudly (never silently shrink, §1.4)."""
    target = f"https://web.archive.org/web/{wayback_ts}id_/{url}" if wayback_ts else url
    cpath = None
    if cache and CACHE_DIR is not None:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cpath = CACHE_DIR / hashlib.sha1(target.encode("utf-8")).hexdigest()
        if cpath.exists() and cpath.stat().st_size > 0:
            return cpath.read_bytes()
    last: Exception | None = None
    for attempt in range(RETRIES):
        try:
            r = requests.get(target, headers=HEADERS, timeout=90)
            if r.status_code == 200 and r.content:
                if cpath is not None:
                    cpath.write_bytes(r.content)
                time.sleep(pause)
                return r.content
            last = RuntimeError(f"HTTP {r.status_code} for {target}")
        except requests.RequestException as exc:  # network flake / Wayback throttle
            last = exc
        time.sleep(5 * (attempt + 1) * (3 if wayback_ts else 1))
    raise RuntimeError(f"failed to fetch {target}: {last}")


def cell(v: Any) -> str:
    if v is None:
        return ""
    s = str(v)
    if s.lower() == "nan":
        return ""
    return re.sub(r"\s+", " ", s.replace("　", " ")).strip()


def tables_from_docx(content: bytes) -> tuple[list[str], list[list[list[str]]]]:
    import docx  # python-docx
    d = docx.Document(io.BytesIO(content))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    tables = [[[cell(c.text) for c in row.cells] for row in t.rows] for t in d.tables]
    return paras, tables


def tables_from_doc(content: bytes, work_dir: Path) -> tuple[list[str], list[list[list[str]]]]:
    """Legacy binary .doc -> tables via Word COM (Hunan 453 precedent)."""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from cn_provincial.extras import parse_doc_wordcom
    work_dir.mkdir(parents=True, exist_ok=True)
    p = work_dir / ("doc_" + hashlib.sha1(content).hexdigest()[:12] + ".doc")
    p.write_bytes(content)
    tables = parse_doc_wordcom(p)
    if not tables:
        raise RuntimeError(f"Word COM returned no tables for {p} (Windows + Word required)")
    return [], [[[cell(c) for c in row] for row in t] for t in tables]


def tables_from_xlsx(content: bytes) -> tuple[list[str], list[list[list[str]]]]:
    import pandas as pd
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None, dtype=str)
    return [], [[[cell(v) for v in row] for row in df.itertuples(index=False)] for df in sheets.values()]


def tables_from_html(content: bytes) -> tuple[list[str], list[list[list[str]]]]:
    import pandas as pd
    text = content.decode("utf-8", errors="replace")
    dfs = pd.read_html(io.StringIO(text), header=None)
    return [], [[[cell(v) for v in row] for row in df.itertuples(index=False)] for df in dfs]


def load_tables(kind: str, content: bytes, work_dir: Path):
    return {"docx": tables_from_docx, "xlsx": tables_from_xlsx, "html": tables_from_html,
            "doc": lambda c: tables_from_doc(c, work_dir)}[kind](content)


# Header-keyword -> field. First match wins per column; order matters.
FIELD_KEYS = [
    ("funder_award_id", ("编号",)),
    ("programme", ("项目名称",)),        # BMSTC: 项目名称 = parent programme when a 课题名称 column exists
    ("title", ("课题名称", "名称", "榜单任务")),
    ("institution", ("依托单位", "承担单位", "申报单位", "揭榜单位", "牵头单位")),
    ("pi", ("申请者", "申请人", "负责人")),
    ("amount_raw", ("金额", "经费")),
    ("discipline", ("学科",)),
    ("scheme", ("项目类型",)),
    ("period", ("起止时间", "起止年月", "实施周期")),
    ("pi_title", ("职称",)),
    ("office", ("主管处室",)),
    ("funding_year", ("资助年度",)),
]


def map_header(row: list[str]) -> dict[str, int] | None:
    """Return {field: column index} if `row` looks like a roster header."""
    joined = "".join(row)
    if "名称" not in joined:
        return None
    out: dict[str, int] = {}
    for i, h in enumerate(row):
        h = h.replace(" ", "")
        if not h or "合作单位" in h:
            continue
        for field, keys in FIELD_KEYS:
            if field in out:
                continue
            if any(k in h for k in keys):
                out[field] = i
                break
    # A lone 项目名称 column (no 课题名称) IS the title.
    if "programme" in out and "title" not in out:
        out["title"] = out.pop("programme")
    return out if "title" in out else None


def section_label(row: list[str]) -> str | None:
    """A merged single-label row inside a table, e.g. '面上项目' / '青年项目'."""
    vals = {c for c in row if c}
    if len(vals) == 1:
        v = next(iter(vals))
        if not re.fullmatch(r"[\d.]+", v) and len(v) <= 30:
            return v
    return None


def rows_from_tables(tables: list[list[list[str]]]) -> list[dict[str, str]]:
    """Walk every table; carry the header map across continuation tables."""
    out: list[dict[str, str]] = []
    hdr: dict[str, int] | None = None
    section: str | None = None
    for t in tables:
        for row in t:
            if not any(row):
                continue
            h = map_header(row)
            if h:
                hdr = h
                continue
            lab = section_label(row)
            if lab:
                section = lab
                continue
            if not hdr:
                continue
            rec = {f: (row[i] if i < len(row) else "") for f, i in hdr.items()}
            if not rec.get("title") or rec["title"] in ("项目名称", "课题名称"):
                continue
            rec["section"] = section or ""
            out.append(rec)
    return out


def synthetic_key(prefix: str, title: str, institution: str) -> str:
    norm = re.sub(r"\s+", "", f"{title}|{institution}").lower()
    return f"{prefix}-" + hashlib.sha1(norm.encode("utf-8")).hexdigest()[:12].upper()


def wan_to_cny(v: str) -> str | None:
    """万元 -> CNY string (×10,000)."""
    m = re.search(r"\d+(?:\.\d+)?", v or "")
    if not m:
        return None
    return f"{float(m.group(0)) * 10000:.0f}"


def parse_period(v: str) -> tuple[str | None, str | None]:
    """'2022-1至2024-12' / '2023年03月至2025年03月' -> (YYYY-MM-01, YYYY-MM-<last>)."""
    nums = re.findall(r"(\d{4})\D{0,2}?(\d{1,2})", v or "")
    if len(nums) < 1:
        return None, None
    def ym(t, end=False):
        y, m = int(t[0]), int(t[1])
        if not 1 <= m <= 12:
            return None
        d = calendar.monthrange(y, m)[1] if end else 1
        return f"{y:04d}-{m:02d}-{d:02d}"
    start = ym(nums[0])
    end = ym(nums[1], end=True) if len(nums) > 1 else None
    return start, end


def first_org(v: str) -> str:
    return re.split(r"[，,、;；]", v or "")[0].strip()


def first_person(v: str) -> str:
    return re.split(r"[，,、;；\s]", (v or "").strip())[0].strip()


def upload_with_shrink_guard(df, parquet_path: Path, bucket: str, key: str, allow_shrink: bool) -> None:
    """Runbook §1.4: never overwrite a bigger corpus with a smaller one."""
    import boto3
    import pandas as pd
    from botocore.exceptions import ClientError
    s3 = boto3.client("s3")
    previous = parquet_path.with_name("_previous_" + parquet_path.name)
    try:
        s3.download_file(bucket, key, str(previous))
        prev_count = len(pd.read_parquet(previous))
        log(f"Shrink check: previous {prev_count}, new {len(df)}")
        if len(df) < prev_count and not allow_shrink:
            raise SystemExit(f"refusing to shrink corpus ({prev_count} -> {len(df)}); "
                             f"rerun with --allow-shrink if genuine")
    except ClientError as e:
        if e.response.get("Error", {}).get("Code") not in {"404", "NoSuchKey", "NotFound"}:
            raise
        log("Shrink check: no existing parquet, first ingest")
    log(f"Uploading to s3://{bucket}/{key}")
    s3.upload_file(str(parquet_path), bucket, key)
    log("Done")
